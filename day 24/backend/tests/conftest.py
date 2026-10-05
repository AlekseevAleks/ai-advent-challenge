"""Общие фикстуры pytest.

Перед каждым тестом подменяются: директория хранения (tmp), клиент эмбеддингов
(FakeEmbeddingsClient). Реальный Ollama не требуется.
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import pytest


@pytest.fixture()
def storage_dir(tmp_path: Path, monkeypatch) -> Path:
    """Изолированное хранилище для каждого теста."""
    target = tmp_path / "storage"
    monkeypatch.setenv("STORAGE_DIR", str(target))
    from app.services import runtime

    runtime.reset_runtime_cache()
    yield target
    runtime.reset_runtime_cache()


@pytest.fixture()
def client(storage_dir):
    """TestClient с подменённым (фейковым) клиентом эмбеддингов."""
    from fastapi.testclient import TestClient

    from app.main import app
    from app.services.ollama_service import ollama_provider

    from tests.fake_embeddings import FakeEmbeddingsClient

    fake = FakeEmbeddingsClient()
    ollama_provider.set_fake(fake)
    with TestClient(app) as c:
        c.fake_embeddings = fake  # type: ignore[attr-defined]
        yield c
    ollama_provider.reset_fake()


@pytest.fixture()
def fake_client():
    from tests.fake_embeddings import FakeEmbeddingsClient

    return FakeEmbeddingsClient()


# ---------------------------------------------------------------------------
# Фабрики тестовых файлов
# ---------------------------------------------------------------------------

def make_pdf(pages: list[str]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    for t in pages:
        pdf.add_page()
        pdf.set_font("helvetica", size=12)
        pdf.cell(text=t)
    return bytes(pdf.output())


def make_docx(paragraphs: list[tuple[int, str]]) -> bytes:
    """paragraphs = [(heading_level|0, text), ...]; level 0 — обычный абзац."""
    import io

    from docx import Document as DocxDocument

    doc = DocxDocument()
    for level, text in paragraphs:
        if level:
            doc.add_heading(text, level=level)
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_md(headings_path: list[tuple[int, str]], body: str) -> str:
    parts = ["# " * lv + h for lv, h in headings_path]
    return "\n".join(parts) + "\n" + body


@pytest.fixture()
def file_factories():
    return {"pdf": make_pdf, "docx": make_docx, "md": make_md}


@pytest.fixture()
def demo_documents(client):
    """Загрузить демонстрационный корпус через API и вернуть документы."""
    r = client.post("/api/documents/demo")
    assert r.status_code == 200, r.text
    return r.json()["documents"]